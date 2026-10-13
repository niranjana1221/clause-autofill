import re, sys
import pdfplumber
from openpyxl import load_workbook
from openpyxl.styles import PatternFill, Alignment

# ------------------------- SETTINGS (edit for a new vehicle) -------------------------
MAIN_TABLE_PAGES = 16
FIRST_DATA_ROW   = 3
COMPONENT_COL    = "A"
COLS = {"B": "C", "D": "E", "F": "G"}
SYNONYMS = {
    "alternator": "generator", "exhuast": "exhaust", "cal": "calibration",
    "ecu": "control unit", "silencer": "exhaust", "scr": "catalytic",
    "ats": "catalytic", "turbo": "pressure charger", "charger": "pressure charger",
}
STOP = {"india","pvt","ltd","limited","private","m/s","option","the","and","with","first",
        "pipe","intermediate","refer","sheet","regular","models","package","type","fan"}
A_STOP = {"engine","system","and","the"}
OVERRIDES = {
    "cal i": {"C": "1.3.1.9.1", "E": "1.3.1.9.3.9", "G": "1.3.1.9.3.9"},
}
MAKE_COLS = {"B": r"make|manufactur|generator", "D": r"type|system", "F": r"type|part|identif"}
# -------------------------------------------------------------------------------------
YEL = PatternFill("solid", fgColor="FFFF00"); RED = PatternFill("solid", fgColor="FF9999")
norm = lambda s: re.sub(r"[^a-z0-9]", "", str(s).lower())

def detect_main_pages(pdf, default=MAIN_TABLE_PAGES):
    for i, page in enumerate(pdf.pages[:60]):
        m = re.search(r"Sheet No\s*:\s*(\d+)\s*[0oO]f\s*(\d+)", page.extract_text() or "")
        if m and m.group(1) == m.group(2):
            return i + 1
    return default

def read_pdf(path, pages=None):
    rows = []
    with pdfplumber.open(path) as pdf:
        n = pages or detect_main_pages(pdf)
        for page in pdf.pages[:n]:
            for table in page.extract_tables():
                for r in table:
                    if len(r) < 3: continue
                    clause, desc, val = [(c or "").replace("\n", " ").strip() for c in r[:3]]
                    if clause.lower().startswith("clause"): continue
                    if clause and re.match(r"^\d", clause):
                        rows.append({"clause": clause.rstrip("."), "desc": desc, "value": val})
                    elif rows:
                        rows[-1]["desc"] += " " + desc; rows[-1]["value"] += " " + val
    descs = {r["clause"]: r["desc"] for r in rows}
    for r in rows:
        parts, ctx = r["clause"].split("."), []
        for i in range(len(parts), 0, -1):
            ctx.append(descs.get(".".join(parts[:i]), ""))
        r["context"] = " ".join(ctx).lower()
        r["nval"] = norm(r["value"]); r["ndesc"] = r["desc"].lower()
    return rows

def part_numbers(t):
    return [norm(p) for p in re.findall(r"[A-Z]{0,3}\s?\d{3}\s\d{3}\s\d{2}\s\d{2}|\b\d[0-9A-Z]{11,16}\b", t)]

def words(t, stop, minlen=4):
    out = []
    for w in re.findall(r"[A-Za-z0-9][A-Za-z0-9+\-]{2,}", t):
        w = SYNONYMS.get(w.lower(), w.lower())
        if w in stop: continue
        if len(w) >= minlen or re.search(r"\d", w) or w in ("dry",): out.append(w)
    return out

def best(rows, text, is_id, comp_words, bonus_re):
    pns = part_numbers(text)
    toks = [(p, 3) for p in pns]
    if not (is_id and pns):
        toks += [(norm(w), 1) for w in words(text, STOP)]
    if not toks: return [], "empty"
    scored = []
    for r in rows:
        s = sum(w for t, w in toks if t in r["nval"])
        if s == 0: continue
        s += 2 * sum(1 for w in comp_words if w in r["context"])
        s += 0.5 if re.search(bonus_re, r["ndesc"]) else 0
        scored.append((s, r["clause"]))
    if not scored: return [], "NOT FOUND"
    scored.sort(key=lambda x: -x[0]); top = scored[0][0]
    if top < 2: return [], "NOT FOUND"
    win = [c for s, c in scored if s == top][:3]
    return win, ("ok" if len(win) == 1 else "check")

def fill(tpl, pdf, pages=None):
    rows = read_pdf(pdf, pages)
    wb = load_workbook(tpl); ws = wb.active
    rep = wb.create_sheet("Report"); rep.append(["Row", "Column", "Clause(s)", "Status"])
    comp = ""
    for r in range(FIRST_DATA_ROW, ws.max_row + 1):
        a = ws[f"{COMPONENT_COL}{r}"].value
        if a: comp = str(a)
        if str(ws[f"F{r}"].value or "").strip().upper() in ("NA", "N/A"):
            for dst in COLS.values():
                c = ws[f"{dst}{r}"]
                if type(c).__name__ != "MergedCell": c.value = "NA"
            continue
        cw = [SYNONYMS.get(w, w) for w in re.findall(r"[a-z]{3,}", comp.lower())]
        cw = [w for w in cw if w not in A_STOP]
        ov = next((v for k, v in OVERRIDES.items() if k in comp.lower()), None)
        for src, dst in COLS.items():
            val, tgt = ws[f"{src}{r}"].value, ws[f"{dst}{r}"]
            if ov and type(tgt).__name__ != "MergedCell" and (val or ws[f"D{r}"].value):
                if dst != "C" or r == FIRST_DATA_ROW or ws[f"A{r}"].value:
                    tgt.value = ov[dst]; tgt.alignment = Alignment(horizontal="center", vertical="center")
                    rep.append([r, f"{src}->{dst}", ov[dst], "override"]); continue
            if not val or type(tgt).__name__ == "MergedCell": continue
            clauses, status = best(rows, str(val), src == "F", cw, MAKE_COLS[src])
            tgt.alignment = Alignment(wrap_text=True, horizontal="center", vertical="center")
            if clauses:
                tgt.value = " / ".join(clauses)
                if status != "ok": tgt.fill = YEL
            else:
                tgt.fill = RED
            rep.append([r, f"{src}->{dst}", tgt.value or "", status])
    return wb, list(rep.iter_rows(min_row=2, values_only=True)), len(rows)

def main(tpl, pdf, out):
    wb, _, n = fill(tpl, pdf)
    wb.save(out); print(f"saved {out} | PDF clauses read: {n}")

if __name__ == "__main__":
    if len(sys.argv) != 4: sys.exit("python autofill_clauses.py template.xlsx table.pdf output.xlsx")
    main(*sys.argv[1:])